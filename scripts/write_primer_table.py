"""Fill the Azenta oligo order form with the top primer pair of each gene.

Usage:
    python write_primer_table.py <results.json> <output.xlsx> [--template <template.xlsx>] [--log <log.csv>]

Each gene with status "ok" gets two rows: <gene>_Fwd and <gene>_Rev.
With --log, also writes a design log (one row per gene: accession, product length,
specificity, status, result URL). The log is kept out of the order form on purpose.
"""

import csv

import argparse
import json
import shutil
from pathlib import Path

import openpyxl

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "azenta_oligo_template.xlsx"
SHEET_NAME = "Azenta Oligo Primer List"
SCALE = "25 nmole"
PURIFICATION = "Desalt"

# Expected header text (lower-cased, prefix match) -> output field
HEADER_KEYS = {
    "oligo #": "number",
    "oligo name": "name",
    "sequence": "seq",
    "length": "length",
    "gc": "gc",
    "tm": "tm",
    "scale": "scale",
    "purification": "purification",
}


def map_columns(ws):
    """Return {field: column_index} by reading the header row."""
    columns = {}
    for cell in ws[1]:
        if cell.value is None:
            continue
        header = str(cell.value).strip().lower()
        for key, field in HEADER_KEYS.items():
            if header.startswith(key) and field not in columns:
                columns[field] = cell.column
                break
    missing = set(HEADER_KEYS.values()) - set(columns)
    if missing:
        raise ValueError(f"Template header is missing columns: {sorted(missing)}")
    return columns


def build_rows(data):
    rows, skipped = [], []
    for job in data.get("jobs", []):
        pair = job.get("pair")
        if job.get("status") != "ok" or not pair:
            skipped.append(f"{job.get('gene')} ({job.get('status')}: {job.get('message', '')})")
            continue
        gene = job["gene"]
        for direction, prefix in (("Fwd", "fwd"), ("Rev", "rev")):
            rows.append({
                "name": f"{gene}_{direction}",
                "seq": pair.get(f"{prefix}_seq"),
                "length": pair.get(f"{prefix}_len"),
                "gc": pair.get(f"{prefix}_gc"),
                "tm": pair.get(f"{prefix}_tm"),
                "scale": SCALE,
                "purification": PURIFICATION,
            })
    return rows, skipped


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("results_json")
    parser.add_argument("output_xlsx")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--log", default=None)
    args = parser.parse_args()

    with open(args.results_json, encoding="utf-8") as fh:
        data = json.load(fh)

    shutil.copyfile(args.template, args.output_xlsx)
    wb = openpyxl.load_workbook(args.output_xlsx)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    columns = map_columns(ws)

    # Clear any existing data rows below the header
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    rows, skipped = build_rows(data)
    for i, row in enumerate(rows, start=1):
        r = i + 1
        ws.cell(row=r, column=columns["number"], value=i)
        for field in ("name", "seq", "length", "gc", "tm", "scale", "purification"):
            ws.cell(row=r, column=columns[field], value=row[field])

    ws.column_dimensions[ws.cell(row=1, column=columns["seq"]).column_letter].width = 32
    ws.column_dimensions[ws.cell(row=1, column=columns["name"]).column_letter].width = 18
    wb.save(args.output_xlsx)

    print(f"Wrote {args.output_xlsx}: {len(rows)} oligos ({len(rows) // 2} genes)")
    for item in skipped:
        print(f"Skipped: {item}")

    if args.log:
        fields = ["gene", "gene_accession", "accession_source", "splice_variants_allowed",
                  "product_min", "product_max", "tm_min", "tm_opt", "tm_max",
                  "status", "message", "product_length", "specificity", "result_url"]
        with open(args.log, "w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=fields)
            writer.writeheader()
            for job in data.get("jobs", []):
                pair = job.get("pair") or {}
                row = {f: job.get(f) for f in fields}
                row["product_length"] = pair.get("product_length")
                row["specificity"] = pair.get("specificity")
                writer.writerow(row)
        print(f"Wrote design log {args.log}")


if __name__ == "__main__":
    main()
